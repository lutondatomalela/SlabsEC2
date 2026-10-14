import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook
from xlsxwriter.exceptions import DuplicateWorksheetName
from slabsec2.xlsx_export import write_workbook


class ExportTests(unittest.TestCase):
    def test_round_trip_and_split(self):
        data = pd.DataFrame({'Panel': [1, 2, 3], 'Value': [1.25, None, -3.5],
                             'Text': ['=literal', 'https://example.org', 'fim']})
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'out.xlsx'
            with patch('slabsec2.xlsx_export.DATA_ROWS_PER_SHEET', 2):
                write_workbook(path, [('Results', data), ('Empty', None)])
            book = load_workbook(path, read_only=True)
            try:
                self.assertEqual(book.sheetnames, ['Results_1', 'Results_2', 'Empty'])
                rows = list(book['Results_1'].values)
                self.assertEqual(rows[1], (1, 1.25, '=literal'))
                self.assertEqual(rows[2], (2, None, 'https://example.org'))
                self.assertEqual(list(book['Results_2'].values)[1], (3, -3.5, 'fim'))
                self.assertEqual(book['Results_1']['C2'].data_type, 's')
            finally:
                book.close()

    def test_full_report_uses_streaming_and_preserves_all_columns(self):
        from slabsec2.gui.reporting import ReportExportMixin

        class Variable:
            def get(self):
                return ""
            def set(self, value):
                pass

        class App(ReportExportMixin):
            def __getattr__(self, name):
                if name.startswith("var_") or name.endswith("_var"):
                    return Variable()
                if name.startswith("df_"):
                    return pd.DataFrame()
                raise AttributeError(name)
            def update_idletasks(self):
                pass

        app = App()
        app.df_results = pd.DataFrame({f"field_{i}": [i, i + 0.5] for i in range(310)})
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsx"
            with patch("tkinter.filedialog.asksaveasfilename", return_value=str(path)), \
                 patch("tkinter.messagebox.showinfo"), \
                 patch("tkinter.messagebox.showerror", side_effect=AssertionError), \
                 patch("pandas.ExcelWriter", side_effect=AssertionError("Non-streaming writer used")):
                app.export_excel()
            book = load_workbook(path, read_only=True)
            try:
                rows = list(book["15_Full_Results"].values)
                self.assertEqual(rows[0], tuple(app.df_results.columns))
                self.assertEqual(rows[1], tuple(range(310)))
                self.assertEqual(rows[2], tuple(i + 0.5 for i in range(310)))
            finally:
                book.close()

    def test_failure_preserves_existing_file(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'out.xlsx'
            path.write_bytes(b'original')
            with self.assertRaises(DuplicateWorksheetName):
                write_workbook(path, [('same', pd.DataFrame({'a': [1]}))] * 2)
            self.assertEqual(path.read_bytes(), b'original')
            self.assertEqual(list(Path(directory).iterdir()), [path])


if __name__ == '__main__':
    unittest.main()
