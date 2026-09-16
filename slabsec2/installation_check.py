"""Verificação da interface e formatos com dados exclusivamente sintéticos."""
from contextlib import ExitStack
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
from unittest.mock import patch


def run_installation_check():
    import pandas as pd
    from openpyxl import load_workbook
    from .gui import SlabsEC2App
    from .runtime import asset_path

    def fail(title, message, **kwargs):
        raise RuntimeError(f"{title}: {message}")

    checks = []
    app = None
    with TemporaryDirectory(prefix="SlabsEC2 verificação ") as tmp, ExitStack() as stack:
        folder = Path(tmp)
        stack.enter_context(patch("tkinter.messagebox.showerror", side_effect=fail))
        stack.enter_context(patch("tkinter.messagebox.showwarning", side_effect=fail))
        stack.enter_context(patch("tkinter.messagebox.showinfo"))
        try:
            app = SlabsEC2App()
            app.report_callback_exception = lambda exc, value, tb: fail("Interface", str(value))
            app.update_idletasks()
            app.update()
            if not app.winfo_exists():
                raise RuntimeError("Interface indisponível")
            checks.append("Interface Tk: OK")
            if not hasattr(app, "_app_icon"):
                raise RuntimeError("Ícone da janela indisponível")
            for name in ("slabsec2.png", "slabsec2.ico"):
                if not asset_path(name).is_file():
                    raise RuntimeError(f"Recurso em falta: {name}")
            checks.append("Ícones PNG/ICO incluídos: OK")
            app.withdraw()
            data = pd.DataFrame([
                {"Panel": 1, "Node": node, "Case": case,
                 "MXX": sign * 22.0 * factor, "MYY": sign * 16.0 * factor,
                 "MXY": 2.0 * factor}
                for node, sign in ((1, 1), (2, -1))
                for case, factor in ((101, 1.0), (203, 0.6), (205, 0.8))
            ])
            for extension in ("csv", "xlsx"):
                path = folder / f"Dados sintéticos.{extension}"
                if extension == "csv":
                    data.to_csv(path, index=False, sep=";", encoding="utf-8-sig")
                else:
                    data.to_excel(path, index=False, engine="openpyxl")
                with patch("tkinter.filedialog.askopenfilename", return_value=str(path)):
                    app.import_file()
                if len(app.df_clean) != len(data):
                    raise RuntimeError(f"Importação incompleta: {extension}")
                checks.append(f"Importação {extension.upper()}: OK")
            app.var_thickness.set("250")
            app.var_cover.set("30")
            app.var_uls_cases.set("101")
            app.var_els_qp_cases.set("203")
            app.var_els_char_cases.set("205")
            app.var_crack_check.set(True)
            app.run_design()
            if len(app.df_results) != len(data) or app.df_summary.empty or app.df_traceability.empty:
                raise RuntimeError("Cálculo ou rastreabilidade incompletos")
            checks.append("Cálculo ELU/ELS e rastreabilidade: OK")
            xlsx = folder / "Resultados sintéticos.xlsx"
            with patch("tkinter.filedialog.asksaveasfilename", return_value=str(xlsx)):
                app.export_excel()
            workbook = load_workbook(xlsx, read_only=True, data_only=True)
            try:
                if len(workbook.sheetnames) < 10:
                    raise RuntimeError("Exportação Excel incompleta")
            finally:
                workbook.close()
            checks.append("Exportação XLSX e reabertura: OK")
            pdf = folder / "Memória sintética.pdf"
            with patch("tkinter.filedialog.asksaveasfilename", return_value=str(pdf)):
                app.export_pdf_report()
            for path in (pdf, pdf.with_name(pdf.stem + "_AUDITORIA.pdf")):
                raw = path.read_bytes()
                if not raw.startswith(b"%PDF-") or b"%%EOF" not in raw[-1024:] or len(raw) < 1000:
                    raise RuntimeError(f"PDF inválido: {path.name}")
            checks.append("Memória de cálculo e auditoria PDF: OK")
        finally:
            if app is not None:
                app.destroy()
    return (f"SlabsEC2 — verificação da instalação\nExecutável autónomo: {bool(getattr(sys, 'frozen', False))}\n"
            + "\n".join(checks))
